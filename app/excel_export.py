# -*- coding: utf-8 -*-
"""Excel 导出：测量结果 + 处理汇总（自动识别 CK 后算防效）+ 参数说明"""
import re
from pathlib import Path
from datetime import datetime
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="2F6FD0")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def _style_header(ws, ncol):
    for c in range(1, ncol+1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

def _autosize(ws, maxw=42):
    for col in ws.columns:
        L = 8
        for c in col:
            if c.value is not None:
                L = max(L, min(maxw, len(str(c.value)) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = L

def _parse(name):
    """从文件名猜 (菌株, 处理)。形如  层出镰刀菌_嘧菌酯_r1 / 水稻纹枯_CK / 3_CK"""
    stem = Path(name).stem
    parts = re.split(r"[_\-\s]+", stem)
    parts = [p for p in parts if p]
    if len(parts) >= 2:
        strain, treat = parts[0], "_".join(parts[:-1]) if re.fullmatch(r"r?\d+", parts[-1]) else "_".join(parts)
        return strain, treat
    return "", stem

def _is_ck(treat):
    return bool(re.search(r"(^|_)ck($|_)", treat, re.I)) or treat.strip().upper() == "CK"

def write_xlsx(path, items, dish_mm, plug_mm):
    wb = openpyxl.Workbook()
    # ---------------- Sheet1 测量结果 ----------------
    ws = wb.active; ws.title = "测量结果"
    hdr = ["序号", "图片文件名", "文件夹", "直径(mm)", "横(mm)", "竖(mm)", "状态", "备注/提示",
           "皿径(mm)", "菌饼(mm)"]
    ws.append(hdr)
    for i, it in enumerate(items, 1):
        mm = it.final_mm() if hasattr(it, "final_mm") else None
        m1, m2 = (it.final_m(), ) if False else (0, 0)
        try: m1, m2 = it.final_m()
        except Exception: m1 = m2 = ""
        ws.append([i, Path(it.path).name, Path(it.path).parent.name,
                   None if mm is None else round(float(mm), 2),
                   None if not m1 else round(float(m1), 2),
                   None if not m2 else round(float(m2), 2),
                   it.status() if hasattr(it, "status") else "失败",
                   (getattr(it, "note", "") or getattr(it, "error", "")),
                   dish_mm, plug_mm])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border = BORDER
        row[3].number_format = "0.00"; row[4].number_format = "0.00"; row[5].number_format = "0.00"
        row[0].alignment = Alignment(horizontal="center")
        row[6].alignment = Alignment(horizontal="center")
    _style_header(ws, len(hdr)); _autosize(ws)

    # ---------------- Sheet2 处理汇总（能识别出菌株/处理时） ----------------
    groups = defaultdict(list); strains = set()
    for it in items:
        mm = it.final_mm() if hasattr(it, "final_mm") else None
        if mm is None: continue
        strain, treat = _parse(Path(it.path).name)
        groups[(strain, treat)].append(float(mm))
        if strain: strains.add(strain)
    made_summary = False
    if len(groups) >= 2 and any(_is_ck(t) for _, t in groups):
        ws2 = wb.create_sheet("处理汇总")
        ws2.append(["菌株/类别", "处理", "重复数", "平均直径(mm)", "CK平均直径(mm)", "防效(%)"])
        ck = {}
        for (strain, treat), vals in groups.items():
            if _is_ck(treat): ck[strain] = sum(vals)/len(vals)
        rows = []
        for (strain, treat), vals in sorted(groups.items()):
            m = sum(vals)/len(vals); c = ck.get(strain)
            eff = None
            if c is not None and not _is_ck(treat) and c > plug_mm:
                eff = (c - m)/(c - plug_mm)*100.0
            ws2.append([strain or "(未识别)", treat, len(vals), round(m, 2),
                        None if c is None else round(c, 2), None if eff is None else round(eff, 2)])
        for row in ws2.iter_rows(min_row=2):
            for c in row: c.border = BORDER
            for idx in (3, 4, 5): row[idx].number_format = "0.00"
            row[2].alignment = Alignment(horizontal="center")
        _style_header(ws2, 6); _autosize(ws2)
        made_summary = True

    # ---------------- Sheet3 参数与说明 ----------------
    ws3 = wb.create_sheet("参数与说明")
    ws3.append(["项目", "值 / 说明"])
    info = [
        ("程序", "皿内菌落直径测量 v1.0"),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("培养皿直径", f"{dish_mm:g} mm"),
        ("菌饼直径", f"{plug_mm:g} mm（黄色枪头 6mm / 蓝色枪头 8mm）"),
        ("测量方法", "以皿径标定像素比例；沿径向采样菌丝纹理找到菌落外缘；过菌落质心作十字交叉，取横、竖两条直径的平均值。也可在界面上用“人工十字测量”点击两点手工量取。"),
        ("防效公式", "防效(%) = (CK平均直径 − 处理平均直径) / (CK平均直径 − 菌饼直径) × 100"),
        ("状态含义", "自动＝未发现异常；需复核＝菌落很小/对比度低/超过皿径等，建议人工确认；人工修正＝已手动设定直径；失败＝未检测到培养皿或识别失败。"),
        ("注意", "毛状菌丝、近满皿、深色培养基等情形识别可能偏小，请重点复核这些图片（状态列会提示）。"),
        ("建议拍摄", "俯拍、皿居中、镜头尽量垂直于皿面、避免强反光与皿底拍摄。"),
    ]
    for k, v in info: ws3.append([k, v])
    for row in ws3.iter_rows(min_row=2):
        row[0].font = Font(bold=True); row[1].alignment = Alignment(wrap_text=True, vertical="top")
    _style_header(ws3, 2)
    ws3.column_dimensions["A"].width = 16; ws3.column_dimensions["B"].width = 96
    if not made_summary:
        ws3.append([]); ws3.append(["提示", "本次文件名未包含「菌株_处理」信息，未生成「处理汇总」页；若文件名形如 层出镰刀菌_嘧菌酯_r1.jpg 则自动生成。"])
    wb.save(path)
    return ws.max_row - 1
